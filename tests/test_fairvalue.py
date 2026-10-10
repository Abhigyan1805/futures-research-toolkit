import math

import pandas as pd
import pytest

from fairvalue import excel, expiry, model, scenarios
from fairvalue.scenarios import Inputs


def test_continuous_known_value():
    assert model.fair_value(100, 0.05, 1.0) == pytest.approx(100 * math.exp(0.05))


def test_discrete_known_value():
    assert model.fair_value(100, 0.05, 2.0, method="discrete") == pytest.approx(110.25)


def test_zero_carry_equals_spot_and_expiry_convergence():
    assert model.fair_value(100, 0.05, 0.0) == 100
    assert model.fair_value(100, 0.05, 0.0, storage=0.01, convenience=0.02) == 100


def test_dividend_lowers_and_storage_raises():
    base = model.fair_value(100, 0.05, 0.5)
    assert model.fair_value(100, 0.05, 0.5, income_yield=0.02) < base
    assert model.fair_value(100, 0.05, 0.5, storage=0.02) > base


def test_implied_roundtrip():
    F = model.fair_value(50, 0.04, 0.3, storage=0.01, income_yield=0.02)
    assert model.implied_income_yield(50, F, 0.04, 0.3, storage=0.01) == pytest.approx(0.02)
    c = model.implied_convenience_yield(50, F, 0.04, 0.3, storage=0.01, income_yield=0.02)
    assert c == pytest.approx(0.0, abs=1e-12)
    Fd = model.fair_value(50, 0.04, 0.3, convenience=0.03, method="discrete")
    c = model.implied_convenience_yield(50, Fd, 0.04, 0.3, method="discrete")
    assert c == pytest.approx(0.03)


def test_backwardation_gives_positive_convenience():
    # futures below spot with positive rates => convenience yield > r + u
    assert model.implied_convenience_yield(100, 95, 0.04, 0.5) > 0.04


def test_vectorised_series():
    s = pd.Series([100.0, 110.0])
    out = model.fair_value(s, 0.05, pd.Series([1.0, 1.0]))
    assert list(out.round(4)) == [round(100 * math.exp(.05), 4), round(110 * math.exp(.05), 4)]


def test_bad_method():
    with pytest.raises(ValueError):
        model.fair_value(1, 0.1, 1, method="x")


def test_shock_and_grid():
    b = Inputs(spot=5000, r=0.04, days=90, income_yield=0.013)
    s = scenarios.shock(b, spot_pct=0.1, rate_bp=100)
    assert s.spot == pytest.approx(5500) and s.r == pytest.approx(0.05)
    g = scenarios.scenario_grid(b, days_list=(90, 30, 0))
    assert len(g) == 5 * 5 * 3
    zero = [r for r in g if r["spot_shock"] == 0 and r["rate_shock_bp"] == 0 and r["days"] == 90][0]
    assert zero["change_vs_base"] == pytest.approx(0)
    at_exp = [r for r in g if r["days"] == 0]
    assert all(r["fair_value"] == pytest.approx(r["spot"]) for r in at_exp)


def test_rate_sensitivity_matches_bump():
    b = Inputs(spot=5000, r=0.04, days=180)
    bump = scenarios.shock(b, rate_bp=1).price() - b.price()
    assert scenarios.rate_sensitivity(b) == pytest.approx(bump, rel=1e-3)


def test_expiry_calendars():
    assert expiry.third_friday(2024, 3).isoformat() == "2024-03-15"
    idx = pd.to_datetime(["2024-03-01", "2024-03-15", "2024-03-18"])
    d = expiry.days_to_front_expiry(idx, "ES")
    assert list(d) == [14, 0, 95]
    cl = expiry.cl_expiries(2024)
    assert cl[1].isoformat() == "2024-01-22"        # Feb-24 contract expired 22 Jan 2024
    assert all(x.weekday() < 5 for x in cl)
    gc = expiry.gc_expiries(2024)
    assert gc[0].isoformat() == "2024-02-27"


def test_excel_formulas(tmp_path):
    from openpyxl import load_workbook
    b = Inputs(spot=2000, r=0.04, days=60, storage=0.005, convenience=0.01)
    p = excel.write_scenario_workbook(tmp_path / "x.xlsx", b, days_list=(60, 30))
    wb = load_workbook(p)
    sc = wb["Scenarios"]
    assert sc.max_row == 1 + 5 * 5 * 2
    assert str(sc["F2"].value).startswith("=D2*EXP(")
    assert wb["Inputs"]["B8"].value.startswith("=")
