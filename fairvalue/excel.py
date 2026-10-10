"""Excel export of the scenario grid, with live formulas.

Sheet "Inputs" holds the base assumptions (change them and everything recalcs);
sheet "Scenarios" has one row per shock combination whose fair value is an
Excel formula  =spot*(1+shock)*EXP((r+bp/10000+storage-yield-conv)*days/365).
"""
from itertools import product

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .scenarios import Inputs

_BLUE = Font(color="0000FF")          # blue = hard-coded input (finance convention)
_HEAD = Font(bold=True)
_FILL = PatternFill("solid", fgColor="DDEBF7")


def write_scenario_workbook(path, base: Inputs, spot_pcts=(-0.10, -0.05, 0.0, 0.05, 0.10),
                            rate_bps=(-100, -50, 0, 50, 100), days_list=None, title="Futures"):
    days_list = tuple(days_list) if days_list is not None else (base.days,)
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    rows = [("Contract", title), ("Spot (S)", base.spot), ("Risk-free rate r (decimal)", base.r),
            ("Storage cost u", base.storage), ("Income/dividend yield q", base.income_yield),
            ("Convenience yield c", base.convenience), ("Days to expiry", base.days),
            ("Base fair value (continuous)", "=B2*EXP((B3+B4-B5-B6)*B7/365)"),
            ("Base fair value (discrete)", "=B2*(1+B3+B4-B5-B6)^(B7/365)")]
    for i, (k, v) in enumerate(rows, 1):
        ws.cell(i, 1, k).font = _HEAD
        c = ws.cell(i, 2, v)
        if i in (2, 3, 4, 5, 6, 7):
            c.font = _BLUE
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    sc = wb.create_sheet("Scenarios")
    heads = ["Spot shock", "Rate shock (bp)", "Days to expiry", "Shocked spot", "Shocked r",
             "Fair value (continuous)", "Change vs base", "Fair value (discrete)"]
    for j, h in enumerate(heads, 1):
        c = sc.cell(1, j, h)
        c.font = _HEAD
        c.fill = _FILL
        sc.column_dimensions[c.column_letter].width = 22
    for i, (sp, rb, d) in enumerate(product(spot_pcts, rate_bps, days_list), 2):
        sc.cell(i, 1, sp).number_format = "0%"
        sc.cell(i, 2, rb)
        sc.cell(i, 3, d)
        for col in (1, 2, 3):
            sc.cell(i, col).font = _BLUE
        sc.cell(i, 4, f"=Inputs!$B$2*(1+A{i})")
        sc.cell(i, 5, f"=Inputs!$B$3+B{i}/10000").number_format = "0.00%"
        carry = f"(E{i}+Inputs!$B$4-Inputs!$B$5-Inputs!$B$6)"
        sc.cell(i, 6, f"=D{i}*EXP({carry}*C{i}/365)").number_format = "#,##0.00"
        sc.cell(i, 7, f"=F{i}-Inputs!$B$8").number_format = "#,##0.00"
        sc.cell(i, 8, f"=D{i}*(1+{carry})^(C{i}/365)").number_format = "#,##0.00"
    sc.freeze_panes = "A2"
    wb.save(path)
    return path
